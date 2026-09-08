from __future__ import annotations

import base64
import html
import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any

import bleach
from bleach.css_sanitizer import CSSSanitizer
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document as WordDocument
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, LETTER, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Image, ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.studio import StudioDocument


ALLOWED_TAGS = {
    "p", "div", "span", "br", "strong", "b", "em", "i", "u", "s", "strike",
    "h1", "h2", "h3", "h4", "blockquote", "ul", "ol", "li", "a", "hr",
    "table", "thead", "tbody", "tfoot", "tr", "th", "td", "img", "sub", "sup",
    "code", "pre",
}
ALLOWED_ATTRIBUTES = {
    "*": [
        "class", "style", "data-page-break", "data-paragraph-style",
        "data-signature-field", "data-field-id", "data-field-type", "data-label",
        "data-assigned-to", "data-required", "data-width", "data-height", "data-placeholder",
    ],
    "a": ["href", "title", "target", "rel"],
    "img": ["src", "alt", "title", "width", "height"],
    "td": ["colspan", "rowspan"],
    "th": ["colspan", "rowspan"],
}
ALLOWED_CSS_PROPERTIES = [
    "color", "background-color", "font-family", "font-size", "font-weight", "font-style",
    "text-decoration", "text-align", "line-height", "letter-spacing", "text-indent",
    "margin-left", "margin-right", "margin-top", "margin-bottom", "padding", "padding-left",
    "padding-right", "padding-top", "padding-bottom", "border", "border-color", "border-width",
    "border-style", "border-radius", "width", "max-width", "min-width", "height", "min-height",
    "max-height", "page-break-after", "page-break-before", "vertical-align", "white-space", "display",
]
CSS_SANITIZER = CSSSanitizer(allowed_css_properties=ALLOWED_CSS_PROPERTIES)


@dataclass(frozen=True)
class Theme:
    accent: str
    heading: str
    soft: str
    body: str
    pdf_font: str
    word_font: str


THEMES: dict[str, Theme] = {
    "guardrisk_orange": Theme("#F47A20", "#3D2A1D", "#FFF4EC", "#25272B", "Helvetica", "Arial"),
    "classic_word": Theme("#2F5597", "#1F3864", "#EAF0F8", "#1F1F1F", "Times-Roman", "Aptos"),
    "classic": Theme("#2F5597", "#1F3864", "#EAF0F8", "#1F1F1F", "Times-Roman", "Aptos"),
    "executive_charcoal": Theme("#30343B", "#202329", "#F0F1F3", "#25272B", "Helvetica", "Arial"),
    "executive": Theme("#30343B", "#202329", "#F0F1F3", "#25272B", "Helvetica", "Arial"),
    "elegant": Theme("#7F6B49", "#4C402D", "#F6F2E9", "#302C25", "Times-Roman", "Georgia"),
    "legal_monochrome": Theme("#111111", "#111111", "#F3F4F6", "#111111", "Times-Roman", "Times New Roman"),
    "legal": Theme("#111111", "#111111", "#F3F4F6", "#111111", "Times-Roman", "Times New Roman"),
    "warm_professional": Theme("#A4552A", "#67351F", "#FFF4EC", "#3F3028", "Helvetica", "Arial"),
    "minimal_clean": Theme("#64748B", "#1E293B", "#F8FAFC", "#1E293B", "Helvetica", "Arial"),
    "clean": Theme("#64748B", "#1E293B", "#F8FAFC", "#1E293B", "Helvetica", "Arial"),
}


def _theme(document: StudioDocument) -> Theme:
    return THEMES.get(document.style_key, THEMES["guardrisk_orange"])


def _settings(document: StudioDocument) -> dict[str, Any]:
    raw = dict(document.settings or {})
    fallback = _number(raw.get("margin_mm"), 20)
    return {
        **raw,
        "page_size": raw.get("page_size", "a4"),
        "orientation": raw.get("orientation", "portrait"),
        "margin_top_mm": _number(raw.get("margin_top_mm"), fallback),
        "margin_right_mm": _number(raw.get("margin_right_mm"), fallback),
        "margin_bottom_mm": _number(raw.get("margin_bottom_mm"), fallback),
        "margin_left_mm": _number(raw.get("margin_left_mm"), fallback),
        "brand_header": raw.get("brand_header", True) is not False,
        "default_font_family": str(raw.get("default_font_family") or "Arial"),
        "default_font_size_pt": max(8, min(36, int(_number(raw.get("default_font_size_pt"), 11)))),
        "default_line_height_percent": max(90, min(250, int(_number(raw.get("default_line_height_percent"), 115)))),
    }


def _number(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def sanitize_document_html(value: str | None) -> str:
    return bleach.clean(
        str(value or "")[:2_500_000],
        tags=ALLOWED_TAGS,
        attributes=ALLOWED_ATTRIBUTES,
        protocols={"http", "https", "mailto", "tel", "data"},
        css_sanitizer=CSS_SANITIZER,
        strip=True,
    )


def plain_text_from_html(value: str | None) -> str:
    soup = BeautifulSoup(value or "", "html.parser")
    for field in soup.select("[data-signature-field]"):
        field.replace_with(f"[{field.get('data-label') or 'Document field'}]")
    return soup.get_text("\n", strip=True)[:1_000_000]


def _style_value(tag: Tag, name: str) -> str | None:
    match = re.search(rf"(?:^|;)\s*{re.escape(name)}\s*:\s*([^;]+)", str(tag.get("style") or ""), re.I)
    return match.group(1).strip() if match else None


def _css_number(value: str | None, default: float = 0.0) -> float:
    if not value:
        return default
    match = re.match(r"\s*(-?[0-9.]+)", value)
    return float(match.group(1)) if match else default


def _font_size_pt(value: str | None, default: float) -> float:
    if not value:
        return default
    size = _css_number(value, default)
    if value.strip().lower().endswith("px"):
        size *= 0.75
    return max(6, min(72, size))


def _data_image(src: str | None) -> bytes | None:
    if not src or not src.startswith("data:image/") or "," not in src:
        return None
    header, encoded = src.split(",", 1)
    if ";base64" not in header:
        return None
    try:
        payload = base64.b64decode(encoded, validate=True)
    except Exception:
        return None
    return payload if len(payload) <= 3 * 1024 * 1024 else None


def _pdf_font(family: str | None, bold: bool = False, italic: bool = False) -> str:
    value = (family or "").lower()
    if "times" in value or "georgia" in value or "serif" in value:
        if bold and italic: return "Times-BoldItalic"
        if bold: return "Times-Bold"
        if italic: return "Times-Italic"
        return "Times-Roman"
    if "courier" in value or "mono" in value:
        if bold and italic: return "Courier-BoldOblique"
        if bold: return "Courier-Bold"
        if italic: return "Courier-Oblique"
        return "Courier"
    if bold and italic: return "Helvetica-BoldOblique"
    if bold: return "Helvetica-Bold"
    if italic: return "Helvetica-Oblique"
    return "Helvetica"


def _pdf_inline(node: Tag | NavigableString, inherited: dict[str, Any] | None = None) -> str:
    attrs = dict(inherited or {})
    if isinstance(node, NavigableString):
        text = html.escape(str(node))
        if not text: return ""
        options: list[str] = []
        if attrs.get("font"): options.append(f'name="{attrs["font"]}"')
        if attrs.get("size"): options.append(f'size="{attrs["size"]}"')
        if attrs.get("color"): options.append(f'color="{attrs["color"]}"')
        return f"<font {' '.join(options)}>{text}</font>" if options else text
    name = node.name.lower()
    if name in {"strong", "b"}: attrs["bold"] = True
    if name in {"em", "i"}: attrs["italic"] = True
    family = _style_value(node, "font-family")
    if family: attrs["font"] = _pdf_font(family, attrs.get("bold", False), attrs.get("italic", False))
    size = _style_value(node, "font-size")
    if size: attrs["size"] = _font_size_pt(size, 10)
    color = _style_value(node, "color")
    if color and re.fullmatch(r"#[0-9a-fA-F]{6}", color): attrs["color"] = color
    inner = "".join(_pdf_inline(child, attrs) for child in node.children)
    if name in {"strong", "b"}: return f"<b>{inner}</b>"
    if name in {"em", "i"}: return f"<i>{inner}</i>"
    if name == "u": return f"<u>{inner}</u>"
    if name == "br": return "<br/>"
    if name in {"sub", "sup"}: return f"<{name}>{inner}</{name}>"
    if name == "code": return f'<font name="Courier">{inner}</font>'
    return inner


def _alignment(tag: Tag) -> int:
    return {"center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY}.get((_style_value(tag, "text-align") or "left").lower(), TA_LEFT)


def _field_info(tag: Tag) -> dict[str, Any]:
    field_type = str(tag.get("data-field-type") or "signature")
    labels = {"signature": "Signature", "initials": "Initials", "date": "Date signed", "name": "Full name", "title": "Title / capacity", "text": "Text field", "checkbox": "Consent"}
    return {
        "field_type": field_type,
        "label": str(tag.get("data-label") or labels.get(field_type, "Document field")),
        "assigned_to": str(tag.get("data-assigned-to") or "Recipient"),
        "required": str(tag.get("data-required") or "false").lower() == "true",
        "placeholder": str(tag.get("data-placeholder") or "Complete here"),
        "height": max(28, min(120, int(_number(tag.get("data-height"), 56)))),
    }


def _pdf_field(tag: Tag, theme: Theme) -> Table:
    info = _field_info(tag)
    required = " · REQUIRED" if info["required"] else ""
    label = Paragraph(f'<b>{html.escape(info["label"])}</b><br/><font size="7" color="#64748B">{html.escape(info["assigned_to"])}{required}</font>', ParagraphStyle("FieldLabel", fontName="Helvetica", fontSize=8.5, leading=10.5, textColor=colors.HexColor(theme.heading)))
    if info["field_type"] == "checkbox":
        value = Paragraph("[ ] &nbsp; " + html.escape(info["placeholder"]), ParagraphStyle("FieldCheck", fontName="Helvetica", fontSize=10, leading=13))
    else:
        lines = "<br/>".join("________________________________________" for _ in range(max(1, min(4, info["height"] // 25))))
        value = Paragraph(lines, ParagraphStyle("FieldValue", fontName="Helvetica", fontSize=9, leading=14, textColor=colors.HexColor("#64748B")))
    table = Table([[label, value]], colWidths=[48 * mm, 105 * mm], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), colors.HexColor(theme.soft)),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(theme.accent)),
        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor(theme.accent)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return table


def _pdf_paragraph(tag: Tag, document: StudioDocument, base: ParagraphStyle, heading_size: float | None = None) -> Paragraph:
    settings = _settings(document)
    size = heading_size or _font_size_pt(_style_value(tag, "font-size"), settings["default_font_size_pt"])
    line_height = _style_value(tag, "line-height")
    if line_height and line_height.endswith("%"):
        leading = size * _css_number(line_height, settings["default_line_height_percent"]) / 100
    elif line_height:
        numeric = _css_number(line_height, 1.15)
        leading = size * numeric if numeric < 4 else numeric
    else:
        leading = size * settings["default_line_height_percent"] / 100
    family = _style_value(tag, "font-family") or settings["default_font_family"]
    style = ParagraphStyle(
        f"StudioParagraph-{id(tag)}",
        parent=base,
        fontName=_pdf_font(family),
        fontSize=size,
        leading=max(size * 1.05, leading),
        alignment=_alignment(tag),
        textColor=colors.HexColor(_theme(document).body),
        leftIndent=_css_number(_style_value(tag, "margin-left")) * mm,
        rightIndent=_css_number(_style_value(tag, "margin-right")) * mm,
        firstLineIndent=_css_number(_style_value(tag, "text-indent")) * mm,
        spaceBefore=_css_number(_style_value(tag, "margin-top"), 0) * 0.75,
        spaceAfter=_css_number(_style_value(tag, "margin-bottom"), 6) * 0.75,
    )
    return Paragraph(_pdf_inline(tag) or " ", style)


def _pdf_table(tag: Tag, document: StudioDocument) -> Table:
    theme = _theme(document)
    base = getSampleStyleSheet()["BodyText"]
    source_rows = tag.find_all("tr")
    width = max((len(row.find_all(["td", "th"], recursive=False)) for row in source_rows), default=1)
    rows: list[list[Any]] = []
    header_rows: set[int] = set()
    for row_index, row in enumerate(source_rows):
        cells = row.find_all(["td", "th"], recursive=False)
        target: list[Any] = []
        for cell in cells:
            if cell.name == "th": header_rows.add(row_index)
            target.append(_pdf_paragraph(cell, document, base))
        target.extend([""] * (width - len(target)))
        rows.append(target)
    table = Table(rows or [[""]], repeatRows=1 if 0 in header_rows else 0, hAlign="LEFT")
    commands: list[tuple[Any, ...]] = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#AEB4BD")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]
    for row_index in header_rows: commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor(theme.soft)))
    table.setStyle(TableStyle(commands))
    return table


def _pdf_story(document: StudioDocument) -> list[Any]:
    soup = BeautifulSoup(sanitize_document_html(document.html_content), "html.parser")
    theme = _theme(document)
    styles = getSampleStyleSheet()
    body = ParagraphStyle("StudioBody", parent=styles["BodyText"], fontName=theme.pdf_font, textColor=colors.HexColor(theme.body))
    story: list[Any] = []
    for node in soup.contents:
        if not isinstance(node, Tag): continue
        if node.get("data-page-break") is not None or _style_value(node, "page-break-after") == "always": story.append(PageBreak())
        elif node.get("data-signature-field") is not None: story.extend([_pdf_field(node, theme), Spacer(1, 3 * mm)])
        elif node.name in {"h1", "h2", "h3", "h4"}:
            size = {"h1": 20, "h2": 15, "h3": 12, "h4": 11}.get(node.name, 11)
            paragraph = _pdf_paragraph(node, document, body, size)
            paragraph.style.textColor = colors.HexColor(theme.heading)
            story.append(paragraph)
        elif node.name in {"p", "div", "blockquote", "pre"}: story.append(_pdf_paragraph(node, document, body))
        elif node.name in {"ul", "ol"}:
            items = [ListItem(_pdf_paragraph(item, document, body), leftIndent=10) for item in node.find_all("li", recursive=False)]
            story.append(ListFlowable(items, bulletType="1" if node.name == "ol" else "bullet", leftIndent=18))
        elif node.name == "table": story.extend([_pdf_table(node, document), Spacer(1, 3 * mm)])
        elif node.name == "img":
            payload = _data_image(str(node.get("src") or ""))
            if payload:
                try: story.append(Image(BytesIO(payload), width=150 * mm, height=95 * mm, kind="proportional"))
                except Exception: pass
        elif node.name == "hr": story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor(theme.accent), spaceBefore=6, spaceAfter=8))
    return story or [Paragraph(" ", body)]


def export_pdf(document: StudioDocument) -> bytes:
    settings = _settings(document)
    page_size = LETTER if settings["page_size"] == "letter" else A4
    if settings["orientation"] == "landscape": page_size = landscape(page_size)
    buffer = BytesIO()
    top = max(settings["margin_top_mm"], 31 if settings["brand_header"] else settings["margin_top_mm"])
    pdf = SimpleDocTemplate(
        buffer, pagesize=page_size,
        rightMargin=settings["margin_right_mm"] * mm, leftMargin=settings["margin_left_mm"] * mm,
        topMargin=top * mm, bottomMargin=max(16, settings["margin_bottom_mm"]) * mm,
        title=document.title, author="gRisk Document Studio", subject="Guardrisk workspace document",
    )
    theme = _theme(document)

    def draw(canvas, doc):
        canvas.saveState()
        width, height = page_size
        if settings["brand_header"]:
            canvas.setFillColor(colors.HexColor(theme.accent)); canvas.setFont("Helvetica-Bold", 17); canvas.drawString(settings["margin_left_mm"] * mm, height - 14 * mm, "G")
            canvas.setFillColor(colors.HexColor(theme.heading)); canvas.setFont("Helvetica-Bold", 9.5); canvas.drawString((settings["margin_left_mm"] + 8) * mm, height - 13.5 * mm, "GUARDRISK")
            canvas.setFont("Helvetica", 6.5); canvas.setFillColor(colors.HexColor("#6B7280")); canvas.drawString((settings["margin_left_mm"] + 8) * mm, height - 17 * mm, "Insurance · Medical Aid · Bonds & Guarantees · Risk Management")
            canvas.setStrokeColor(colors.HexColor(theme.accent)); canvas.setLineWidth(1.2); canvas.line(settings["margin_left_mm"] * mm, height - 20 * mm, width - settings["margin_right_mm"] * mm, height - 20 * mm)
        canvas.setFont("Helvetica", 7); canvas.setFillColor(colors.HexColor("#7A8088")); canvas.drawRightString(width - settings["margin_right_mm"] * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    pdf.build(_pdf_story(document), onFirstPage=draw, onLaterPages=draw)
    return buffer.getvalue()


def _word_font(run, family: str | None) -> None:
    if not family: return
    run.font.name = family
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts"); r_pr.insert(0, r_fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"): r_fonts.set(qn(f"w:{key}"), family)


def _word_inline(paragraph, node: Tag | NavigableString, inherited: dict[str, Any] | None = None) -> None:
    attrs = dict(inherited or {})
    if isinstance(node, NavigableString):
        if str(node):
            run = paragraph.add_run(str(node)); run.bold = attrs.get("bold", False); run.italic = attrs.get("italic", False); run.underline = attrs.get("underline", False); run.font.strike = attrs.get("strike", False); run.font.subscript = attrs.get("subscript", False); run.font.superscript = attrs.get("superscript", False)
            _word_font(run, attrs.get("font_family"))
            if attrs.get("size"): run.font.size = Pt(attrs["size"])
            if attrs.get("color"):
                try: run.font.color.rgb = RGBColor.from_string(attrs["color"].lstrip("#"))
                except Exception: pass
        return
    name = node.name.lower()
    if name in {"strong", "b"}: attrs["bold"] = True
    if name in {"em", "i"}: attrs["italic"] = True
    if name == "u": attrs["underline"] = True
    if name in {"s", "strike"}: attrs["strike"] = True
    if name == "sub": attrs["subscript"] = True
    if name == "sup": attrs["superscript"] = True
    family = _style_value(node, "font-family")
    if family: attrs["font_family"] = family.split(",")[0].strip(" '\"")
    size = _style_value(node, "font-size")
    if size: attrs["size"] = _font_size_pt(size, 11)
    color = _style_value(node, "color")
    if color and re.fullmatch(r"#[0-9a-fA-F]{6}", color): attrs["color"] = color
    if name == "br": paragraph.add_run().add_break(); return
    for child in node.children: _word_inline(paragraph, child, attrs)


def _word_paragraph(container, tag: Tag, document: StudioDocument, style: str | None = None):
    settings = _settings(document)
    paragraph = container.add_paragraph(style=style)
    paragraph.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get((_style_value(tag, "text-align") or "left").lower(), WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.line_spacing = settings["default_line_height_percent"] / 100
    paragraph.paragraph_format.space_before = Pt(_css_number(_style_value(tag, "margin-top"), 0) * 0.75)
    paragraph.paragraph_format.space_after = Pt(_css_number(_style_value(tag, "margin-bottom"), 6) * 0.75)
    paragraph.paragraph_format.left_indent = Mm(_css_number(_style_value(tag, "margin-left"), 0))
    paragraph.paragraph_format.right_indent = Mm(_css_number(_style_value(tag, "margin-right"), 0))
    paragraph.paragraph_format.first_line_indent = Mm(_css_number(_style_value(tag, "text-indent"), 0))
    _word_inline(paragraph, tag, {"font_family": _style_value(tag, "font-family") or settings["default_font_family"], "size": _font_size_pt(_style_value(tag, "font-size"), settings["default_font_size_pt"])})
    return paragraph


def _cell_border(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None: borders = OxmlElement("w:tcBorders"); tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None: node = OxmlElement(f"w:{edge}"); borders.append(node)
        node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "7"); node.set(qn("w:color"), color.lstrip("#"))


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr(); node = tc_pr.find(qn("w:shd"))
    if node is None: node = OxmlElement("w:shd"); tc_pr.append(node)
    node.set(qn("w:fill"), fill.lstrip("#"))


def _word_field(word: WordDocument, tag: Tag, document: StudioDocument) -> None:
    theme = _theme(document); info = _field_info(tag)
    table = word.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    table.columns[0].width = Mm(48); table.columns[1].width = Mm(105)
    label, value = table.rows[0].cells; _shade(label, theme.soft)
    for cell in (label, value): _cell_border(cell, theme.accent); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    label.text = info["label"]; label.paragraphs[0].runs[0].bold = True; label.paragraphs[0].runs[0].font.size = Pt(9)
    meta = label.add_paragraph(f'{info["assigned_to"]}{" · REQUIRED" if info["required"] else ""}'); meta.runs[0].font.size = Pt(7); meta.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    value.text = f'☐  {info["placeholder"]}' if info["field_type"] == "checkbox" else "\n".join("________________________________" for _ in range(max(1, min(4, info["height"] // 25))))
    value.paragraphs[0].runs[0].font.size = Pt(9)
    word.add_paragraph().paragraph_format.space_after = Pt(0)


def _word_content(word: WordDocument, document: StudioDocument) -> None:
    soup = BeautifulSoup(sanitize_document_html(document.html_content), "html.parser")
    settings = _settings(document)
    for node in soup.contents:
        if not isinstance(node, Tag): continue
        if node.get("data-page-break") is not None or _style_value(node, "page-break-after") == "always": word.add_page_break()
        elif node.get("data-signature-field") is not None: _word_field(word, node, document)
        elif node.name in {"h1", "h2", "h3", "h4"}: _word_paragraph(word, node, document, f"Heading {min(int(node.name[1]), 3)}")
        elif node.name in {"p", "div", "blockquote", "pre"}: _word_paragraph(word, node, document)
        elif node.name in {"ul", "ol"}:
            style = "List Number" if node.name == "ol" else "List Bullet"
            for item in node.find_all("li", recursive=False): _word_paragraph(word, item, document, style)
        elif node.name == "table":
            rows = node.find_all("tr"); width = max((len(row.find_all(["td", "th"], recursive=False)) for row in rows), default=1); table = word.add_table(rows=0, cols=width); table.style = "Table Grid"
            for source_row in rows:
                target = table.add_row().cells
                for index, cell in enumerate(source_row.find_all(["td", "th"], recursive=False)):
                    target[index].text = ""
                    if cell.name == "th": _shade(target[index], _theme(document).soft)
                    _word_inline(target[index].paragraphs[0], cell, {"bold": cell.name == "th", "font_family": settings["default_font_family"], "size": settings["default_font_size_pt"] - 1})
        elif node.name == "img":
            payload = _data_image(str(node.get("src") or ""))
            if payload:
                try: word.add_picture(BytesIO(payload), width=Mm(150))
                except Exception: pass
        elif node.name == "hr":
            paragraph = word.add_paragraph(); p_pr = paragraph._p.get_or_add_pPr(); borders = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom"); bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "8"); bottom.set(qn("w:color"), _theme(document).accent.lstrip("#")); borders.append(bottom); p_pr.append(borders)


def export_docx(document: StudioDocument) -> bytes:
    settings = _settings(document); theme = _theme(document); word = WordDocument(); section = word.sections[0]
    if settings["page_size"] == "letter": section.page_width = Inches(8.5); section.page_height = Inches(11)
    else: section.page_width = Mm(210); section.page_height = Mm(297)
    if settings["orientation"] == "landscape": section.orientation = WD_ORIENT.LANDSCAPE; section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Mm(max(31, settings["margin_top_mm"] if settings["brand_header"] else settings["margin_top_mm"])); section.right_margin = Mm(settings["margin_right_mm"]); section.bottom_margin = Mm(settings["margin_bottom_mm"]); section.left_margin = Mm(settings["margin_left_mm"])

    normal = word.styles["Normal"]; normal.font.name = settings["default_font_family"]; normal.font.size = Pt(settings["default_font_size_pt"]); normal.font.color.rgb = RGBColor.from_string(theme.body.lstrip("#")); normal.paragraph_format.line_spacing = settings["default_line_height_percent"] / 100; normal.paragraph_format.space_after = Pt(6)
    for index, size in ((1, 20), (2, 15), (3, 12)):
        style = word.styles[f"Heading {index}"]; style.font.name = settings["default_font_family"]; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor.from_string(theme.heading.lstrip("#"))

    if settings["brand_header"]:
        header = section.header; table = header.add_table(rows=1, cols=2, width=section.page_width - section.left_margin - section.right_margin); table.alignment = WD_TABLE_ALIGNMENT.CENTER
        left = table.cell(0, 0).paragraphs[0]; run = left.add_run("G  GUARDRISK"); run.bold = True; run.font.size = Pt(12); run.font.color.rgb = RGBColor.from_string(theme.accent.lstrip("#"))
        right = table.cell(0, 1).paragraphs[0]; right.alignment = WD_ALIGN_PARAGRAPH.RIGHT; detail = right.add_run("Insurance · Medical Aid\nBonds & Guarantees · Risk Management"); detail.font.size = Pt(6.5); detail.font.color.rgb = RGBColor(100, 116, 139)
        border = header.add_paragraph(); p_pr = border._p.get_or_add_pPr(); p_bdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom"); bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12"); bottom.set(qn("w:color"), theme.accent.lstrip("#")); p_bdr.append(bottom); p_pr.append(p_bdr)

    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; footer.add_run(f"gRisk · Version {document.version} · Page ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    core = word.core_properties; core.title = document.title; core.author = "gRisk Document Studio"; core.subject = "Guardrisk workspace document"; core.comments = "Generated by gRisk Document Studio"
    _word_content(word, document)
    buffer = BytesIO(); word.save(buffer); return buffer.getvalue()
