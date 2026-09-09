from __future__ import annotations

from datetime import datetime
from io import BytesIO
from types import SimpleNamespace
from typing import Any
from zoneinfo import ZoneInfo

from docx import Document as WordDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as pdf_canvas

from app.models.studio import StudioDocument
from app.services.studio_export import export_docx as export_base_docx
from app.services.studio_export import export_pdf as export_base_pdf

ORANGE = "#F15A24"
PEACH = "#F7B38C"
CHARCOAL = "#344054"
MUTED = "#667085"
LINE = "#D0D5DD"

DEFAULT_ADDRESS = "LNDC Centre, Ground Floor, Shop No. ____"
DEFAULT_PHONE_1 = "+266 2232 2537"
DEFAULT_PHONE_2 = "+266 6272 0488"
DEFAULT_EMAIL = "info@guardrisk.co.ls"
DEFAULT_FOOTER_LEFT = "GUARDRISK HEALTH"
DEFAULT_FOOTER_RIGHT = "LOW COST MEDICAL AID"
DEFAULT_RECIPIENT = "Recipient Name"
DEFAULT_COMPANY = "Company Name"
DEFAULT_SUBJECT = "Your Subject Line Goes Here"
DEFAULT_TAGLINE = "YOUR LINK TO PREMIER HEALTHCARE"


def _number(value: Any, default: float) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def normalize_lesotho_phone(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    digits = "".join(character for character in text if character.isdigit())
    if digits.startswith("266"):
        local = digits[3:]
    elif digits.startswith("0"):
        local = digits[1:]
    else:
        local = digits
    if not local:
        return "+266"
    grouped = " ".join(local[index : index + 4] for index in range(0, len(local), 4))
    return f"+266 {grouped}"


def _maseru_now() -> str:
    return datetime.now(ZoneInfo("Africa/Maseru")).strftime("%d %B %Y, %H:%M")


def _document_code(document: StudioDocument | None) -> str:
    if document is None or getattr(document, "id", None) is None:
        return "GRK-00000"
    tail = str(document.id).replace("-", "")[-8:]
    try:
        number = int(tail, 16) % 100000
    except ValueError:
        number = 0
    return f"GRK-{number:05d}"


def stationery_values(
    settings: dict[str, Any] | None,
    document: StudioDocument | None = None,
) -> dict[str, str]:
    source = settings or {}
    return {
        "date_time": str(source.get("letterhead_date_time") or _maseru_now()),
        "code": str(source.get("letterhead_code") or _document_code(document)),
        "recipient": str(source.get("letterhead_recipient") or DEFAULT_RECIPIENT),
        "company": str(source.get("letterhead_company") or DEFAULT_COMPANY),
        "address": str(source.get("letterhead_address") or DEFAULT_ADDRESS),
        "subject": str(source.get("letterhead_subject") or DEFAULT_SUBJECT),
        "tagline": str(source.get("letterhead_tagline") or DEFAULT_TAGLINE),
        "phone_1": normalize_lesotho_phone(source.get("letterhead_phone_1") or DEFAULT_PHONE_1),
        "phone_2": normalize_lesotho_phone(source.get("letterhead_phone_2") or DEFAULT_PHONE_2),
        "email": str(source.get("letterhead_email") or DEFAULT_EMAIL),
        "footer_left": str(source.get("letterhead_footer_left") or DEFAULT_FOOTER_LEFT),
        "footer_right": str(source.get("letterhead_footer_right") or DEFAULT_FOOTER_RIGHT),
    }


def official_settings(
    value: dict[str, Any] | None,
    document: StudioDocument | None = None,
) -> dict[str, Any]:
    settings = dict(value or {})
    settings["brand_header"] = True
    settings["official_letterhead"] = "guardrisk_editable_v4"
    details = stationery_values(settings, document)
    for key, detail in details.items():
        settings[f"letterhead_{key}"] = detail
    return settings


def _export_proxy(document: StudioDocument) -> SimpleNamespace:
    settings = official_settings(document.settings, document)
    settings["brand_header"] = False
    settings["margin_top_mm"] = max(96.0, _number(settings.get("margin_top_mm"), 20.0) + 76.0)
    settings["margin_bottom_mm"] = max(40.0, _number(settings.get("margin_bottom_mm"), 20.0) + 20.0)
    settings["margin_left_mm"] = max(18.0, _number(settings.get("margin_left_mm"), 20.0))
    settings["margin_right_mm"] = max(18.0, _number(settings.get("margin_right_mm"), 20.0))
    return SimpleNamespace(
        title=document.title,
        style_key=document.style_key,
        html_content=document.html_content,
        settings=settings,
        version=document.version,
    )


def _draw_label_value(
    canvas: pdf_canvas.Canvas,
    *,
    label: str,
    value: str,
    x: float,
    y: float,
    label_width: float = 18 * mm,
) -> None:
    canvas.setFillColor(colors.HexColor(CHARCOAL))
    canvas.setFont("Helvetica-Bold", 6.7)
    canvas.drawString(x, y, label)
    canvas.setFont("Helvetica", 6.7)
    canvas.drawString(x + label_width, y, value)


def _draw_pdf_stationery(
    canvas: pdf_canvas.Canvas,
    width: float,
    height: float,
    details: dict[str, str],
) -> None:
    orange = colors.HexColor(ORANGE)
    peach = colors.HexColor(PEACH)
    charcoal = colors.HexColor(CHARCOAL)
    muted = colors.HexColor(MUTED)
    line = colors.HexColor(LINE)

    center = width / 2
    canvas.setFillColor(orange)
    canvas.circle(center, height - 12 * mm, 8.2 * mm, stroke=0, fill=1)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 20)
    canvas.drawCentredString(center, height - 15.1 * mm, "G")
    canvas.setFillColor(orange)
    canvas.setFont("Times-Bold", 20)
    canvas.drawCentredString(center, height - 29 * mm, "GUARDRISK")

    canvas.setStrokeColor(orange)
    canvas.setLineWidth(0.8)
    canvas.line(18 * mm, height - 36 * mm, 57 * mm, height - 36 * mm)
    canvas.line(width - 57 * mm, height - 36 * mm, width - 18 * mm, height - 36 * mm)
    canvas.setFillColor(charcoal)
    canvas.setFont("Helvetica", 6.5)
    canvas.drawCentredString(center, height - 37.3 * mm, details["tagline"].upper())

    start_y = height - 49 * mm
    left = 18 * mm
    _draw_label_value(canvas, label="Date:", value=details["date_time"], x=left, y=start_y)
    _draw_label_value(canvas, label="Code:", value=details["code"], x=left, y=start_y - 5 * mm)
    _draw_label_value(canvas, label="To:", value=details["recipient"], x=left, y=start_y - 10 * mm)
    _draw_label_value(canvas, label="Company:", value=details["company"], x=left, y=start_y - 15 * mm)

    divider_x = center
    canvas.setStrokeColor(line)
    canvas.setLineWidth(0.5)
    canvas.line(divider_x, height - 45 * mm, divider_x, height - 68 * mm)

    right_x = center + 11 * mm
    canvas.setFillColor(charcoal)
    canvas.setFont("Helvetica-Bold", 6.7)
    canvas.drawString(right_x, start_y, "Address:")
    canvas.setFont("Helvetica", 6.7)
    address_width = 37
    address_lines = [details["address"][index : index + address_width] for index in range(0, len(details["address"]), address_width)] or [""]
    for index, chunk in enumerate(address_lines[:4]):
        canvas.drawString(right_x + 20 * mm, start_y - index * 4 * mm, chunk)

    subject_y = height - 79 * mm
    canvas.setFillColor(orange)
    canvas.setFont("Helvetica-Bold", 11)
    canvas.drawString(18 * mm, subject_y, "RE:")
    canvas.drawString(31 * mm, subject_y, details["subject"])
    canvas.setStrokeColor(orange)
    canvas.setLineWidth(0.7)
    canvas.line(18 * mm, subject_y - 3 * mm, width - 18 * mm, subject_y - 3 * mm)

    footer_top = 34 * mm
    canvas.setFillColor(colors.white)
    canvas.rect(0, 0, width, footer_top, stroke=0, fill=1)
    canvas.setStrokeColor(line)
    canvas.setLineWidth(0.45)
    canvas.line(18 * mm, 30 * mm, width - 18 * mm, 30 * mm)
    canvas.setFillColor(muted)
    canvas.setFont("Helvetica", 5.8)
    canvas.drawString(18 * mm, 24 * mm, f"Address: {details['address']}")
    canvas.drawCentredString(center, 24 * mm, f"Call: {details['phone_1']} / {details['phone_2']}")
    canvas.drawRightString(width - 18 * mm, 24 * mm, f"Email: {details['email']}")

    canvas.setFillColor(orange)
    wave = canvas.beginPath()
    wave.moveTo(0, 12 * mm)
    wave.curveTo(34 * mm, 3 * mm, 64 * mm, 9 * mm, 94 * mm, 12 * mm)
    wave.curveTo(125 * mm, 15 * mm, 148 * mm, 5 * mm, 177 * mm, 7 * mm)
    wave.curveTo(194 * mm, 8 * mm, 204 * mm, 12 * mm, width, 15 * mm)
    wave.lineTo(width, 0)
    wave.lineTo(0, 0)
    wave.close()
    canvas.drawPath(wave, stroke=0, fill=1)
    canvas.setStrokeColor(peach)
    canvas.setLineWidth(2)
    canvas.bezier(0, 14 * mm, 34 * mm, 5 * mm, 64 * mm, 11 * mm, 94 * mm, 14 * mm)
    canvas.bezier(94 * mm, 14 * mm, 126 * mm, 17 * mm, 150 * mm, 7 * mm, width, 17 * mm)

    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 5.6)
    canvas.drawString(10 * mm, 5.5 * mm, details["footer_left"].upper())
    canvas.drawRightString(width - 10 * mm, 5.5 * mm, details["footer_right"].upper())


def export_pdf(document: StudioDocument) -> bytes:
    source = export_base_pdf(_export_proxy(document))
    details = stationery_values(document.settings, document)
    reader = PdfReader(BytesIO(source))
    writer = PdfWriter()
    for page in reader.pages:
        width = float(page.mediabox.width)
        height = float(page.mediabox.height)
        overlay_buffer = BytesIO()
        overlay_canvas = pdf_canvas.Canvas(overlay_buffer, pagesize=(width, height))
        _draw_pdf_stationery(overlay_canvas, width, height, details)
        overlay_canvas.save()
        overlay_buffer.seek(0)
        page.merge_page(PdfReader(overlay_buffer).pages[0], over=True)
        writer.add_page(page)
    if reader.metadata:
        writer.add_metadata({key: str(value) for key, value in reader.metadata.items() if value is not None})
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _set_run(run: Any, *, size: float, bold: bool = False, color: str = CHARCOAL) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))


def _top_border(cell: Any, color: str = LINE) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    top = borders.find(qn("w:top"))
    if top is None:
        top = OxmlElement("w:top")
        borders.append(top)
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "4")
    top.set(qn("w:color"), color.lstrip("#"))


def _bottom_border(cell: Any, color: str = ORANGE) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), color.lstrip("#"))


def _clear_story(container: Any) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)


def _add_docx_header(section: Any, details: dict[str, str]) -> None:
    header = section.header
    header.is_linked_to_previous = False
    _clear_story(header)

    logo = header.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_after = Pt(0)
    _set_run(logo.add_run("G"), size=18, bold=True, color=ORANGE)

    brand = header.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.paragraph_format.space_after = Pt(0)
    _set_run(brand.add_run("GUARDRISK"), size=17, bold=True, color=ORANGE)

    tagline = header.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline.paragraph_format.space_after = Pt(5)
    _set_run(tagline.add_run(details["tagline"].upper()), size=6.2, color=CHARCOAL)

    available = section.page_width - section.left_margin - section.right_margin
    meta = header.add_table(rows=1, cols=2, width=available)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    left_cell, right_cell = meta.rows[0].cells
    left_lines = (
        ("Date: ", details["date_time"]),
        ("Code: ", details["code"]),
        ("To: ", details["recipient"]),
        ("Company: ", details["company"]),
    )
    left_paragraph = left_cell.paragraphs[0]
    left_paragraph.paragraph_format.space_after = Pt(0)
    for index, (label, value) in enumerate(left_lines):
        _set_run(left_paragraph.add_run(label), size=6.3, bold=True)
        _set_run(left_paragraph.add_run(value), size=6.3, color=MUTED)
        if index < len(left_lines) - 1:
            left_paragraph.add_run().add_break()

    address = right_cell.paragraphs[0]
    address.alignment = WD_ALIGN_PARAGRAPH.LEFT
    address.paragraph_format.space_after = Pt(0)
    _set_run(address.add_run("Address: "), size=6.3, bold=True)
    _set_run(address.add_run(details["address"]), size=6.3, color=MUTED)

    subject = header.add_table(rows=1, cols=1, width=available)
    subject.alignment = WD_TABLE_ALIGNMENT.CENTER
    subject_cell = subject.cell(0, 0)
    _bottom_border(subject_cell)
    subject_paragraph = subject_cell.paragraphs[0]
    subject_paragraph.paragraph_format.space_before = Pt(4)
    subject_paragraph.paragraph_format.space_after = Pt(2)
    _set_run(subject_paragraph.add_run("RE: "), size=10.5, bold=True, color=ORANGE)
    _set_run(subject_paragraph.add_run(details["subject"]), size=10.5, bold=True, color=ORANGE)


def _add_docx_footer(section: Any, details: dict[str, str]) -> None:
    footer = section.footer
    footer.is_linked_to_previous = False
    _clear_story(footer)
    available = section.page_width - section.left_margin - section.right_margin
    table = footer.add_table(rows=1, cols=3, width=available)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    values = (
        f"Address: {details['address']}",
        f"Call: {details['phone_1']} / {details['phone_2']}",
        f"Email: {details['email']}",
    )
    for cell, value in zip(table.rows[0].cells, values, strict=True):
        _top_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        _set_run(paragraph.add_run(value), size=5.6, color=MUTED)

    labels = footer.add_paragraph()
    labels.paragraph_format.space_before = Pt(2)
    labels.paragraph_format.space_after = Pt(0)
    _set_run(labels.add_run(details["footer_left"].upper()), size=5.5, bold=True, color=ORANGE)
    labels.add_run(" " * 20)
    _set_run(labels.add_run(details["footer_right"].upper()), size=5.5, bold=True, color=ORANGE)


def export_docx(document: StudioDocument) -> bytes:
    source = export_base_docx(_export_proxy(document))
    details = stationery_values(document.settings, document)
    word = WordDocument(BytesIO(source))
    for section in word.sections:
        section.top_margin = Mm(96)
        section.bottom_margin = Mm(40)
        section.left_margin = Mm(max(18.0, _number((document.settings or {}).get("margin_left_mm"), 20.0)))
        section.right_margin = Mm(max(18.0, _number((document.settings or {}).get("margin_right_mm"), 20.0)))
        section.header_distance = Mm(3)
        section.footer_distance = Mm(3)
        _add_docx_header(section, details)
        _add_docx_footer(section, details)

    core = word.core_properties
    core.author = "Guardrisk Insurance Brokers"
    core.subject = "Official Guardrisk correspondence"
    core.comments = "Generated by gRisk Document Studio using editable Guardrisk stationery"
    output = BytesIO()
    word.save(output)
    return output.getvalue()
