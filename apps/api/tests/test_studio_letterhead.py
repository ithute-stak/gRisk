import uuid
from datetime import datetime, timezone
from io import BytesIO

from docx import Document as WordDocument
from pypdf import PdfReader

from app.models.studio import StudioDocument
from app.services.studio_letterhead import (
    export_docx,
    export_pdf,
    normalize_lesotho_phone,
    official_settings,
    stationery_values,
)
from app.services.studio_system_stamp import build_system_stamp, render_system_stamp_png

CUSTOM_DATE = "09 September 2026, 11:16"
CUSTOM_RECIPIENT = "Mpho Mosotho"
CUSTOM_COMPANY = "Example Company"
CUSTOM_SUBJECT = "Medical Aid Confirmation"
CUSTOM_TAGLINE = "Your Link to Premier Healthcare"
CUSTOM_ADDRESS_1 = "LNDC Centre, Ground Floor,"
CUSTOM_ADDRESS_2 = "Shop No. 12,"
CUSTOM_ADDRESS_3 = "Maseru 100, Lesotho"
CUSTOM_CLOSING_1 = "Kind regards,"
CUSTOM_CLOSING_2 = "For and on behalf of Guardrisk."
CUSTOM_SIGNER = "M. Mosotho"
CUSTOM_TITLE = "Authorised Signatory"
CUSTOM_SIGNATURE = "Click here to digitally sign"
USER_CONTROLLED_STAMP = "USER CONTROLLED STAMP"
DOCUMENT_ID = uuid.UUID("12345678-1234-5678-1234-56781234abcd")


def _document() -> StudioDocument:
    html = """
    <p>Dear Mpho Mosotho,</p>
    <p>First page body content.</p>
    <div data-page-break="true"></div>
    <p>Second page body content.</p>
    """
    return StudioDocument(
        id=DOCUMENT_ID,
        owner_user_id=uuid.uuid4(),
        title="Letterhead Test",
        template_key="formal_letter",
        style_key="classic_word",
        content_json={"type": "doc", "content": []},
        html_content=html,
        plain_text="Dear Mpho Mosotho First page body content. Second page body content.",
        settings={
            "page_size": "a4",
            "orientation": "portrait",
            "margin_top_mm": 12,
            "margin_right_mm": 12,
            "margin_bottom_mm": 12,
            "margin_left_mm": 12,
            "brand_header": False,
            "letterhead_date_time": CUSTOM_DATE,
            "letterhead_code": "GRK-LEGACY",
            "letterhead_recipient": CUSTOM_RECIPIENT,
            "letterhead_company": CUSTOM_COMPANY,
            "letterhead_address_line_1": CUSTOM_ADDRESS_1,
            "letterhead_address_line_2": CUSTOM_ADDRESS_2,
            "letterhead_address_line_3": CUSTOM_ADDRESS_3,
            "letterhead_subject": CUSTOM_SUBJECT,
            "letterhead_tagline": CUSTOM_TAGLINE,
            "letterhead_phone_1": "22322537",
            "letterhead_phone_2": "062720488",
            "letterhead_email": "custom@guardrisk.co.ls",
            "letterhead_footer_left": "Guardrisk Health",
            "letterhead_footer_right": "Low Cost Medical Aid",
            "letterhead_closing_line_1": CUSTOM_CLOSING_1,
            "letterhead_closing_line_2": CUSTOM_CLOSING_2,
            "letterhead_signer_name": CUSTOM_SIGNER,
            "letterhead_signer_title": CUSTOM_TITLE,
            "letterhead_signature_label": CUSTOM_SIGNATURE,
            "letterhead_stamp_label": USER_CONTROLLED_STAMP,
        },
        version=4,
    )


def _container_text(container) -> str:
    paragraphs = [paragraph.text for paragraph in container.paragraphs]
    tables = [cell.text for table in container.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *tables])


def _document_body_text(word: WordDocument) -> str:
    return "\n".join(word.element.body.itertext())


def test_phone_normalization_always_adds_lesotho_country_code() -> None:
    assert normalize_lesotho_phone("22322537") == "+266 2232 2537"
    assert normalize_lesotho_phone("062720488") == "+266 6272 0488"
    assert normalize_lesotho_phone("+266 5939 5332") == "+266 5939 5332"


def test_official_settings_remove_user_controlled_stamp_and_legacy_code() -> None:
    document = _document()
    settings = official_settings(document.settings, document)
    assert settings["brand_header"] is True
    assert settings["official_letterhead"] == "guardrisk_reference_v7_backend_stamp"
    assert "letterhead_code" not in settings
    assert "letterhead_stamp_label" not in settings
    assert settings["letterhead_date_time"] == CUSTOM_DATE
    assert settings["letterhead_recipient"] == CUSTOM_RECIPIENT
    assert settings["letterhead_company"] == CUSTOM_COMPANY
    assert settings["letterhead_subject"] == CUSTOM_SUBJECT
    assert settings["letterhead_tagline"] == CUSTOM_TAGLINE
    assert settings["letterhead_address_line_1"] == CUSTOM_ADDRESS_1
    assert settings["letterhead_address_line_2"] == CUSTOM_ADDRESS_2
    assert settings["letterhead_address_line_3"] == CUSTOM_ADDRESS_3
    assert settings["letterhead_phone_1"] == "+266 2232 2537"
    assert settings["letterhead_phone_2"] == "+266 6272 0488"
    assert settings["letterhead_signer_name"] == CUSTOM_SIGNER


def test_stationery_values_have_no_code_or_stamp_and_keep_three_address_lines() -> None:
    details = stationery_values(_document().settings)
    assert "code" not in details
    assert "stamp_label" not in details
    assert details["date_time"] == CUSTOM_DATE
    assert details["address_line_1"] == CUSTOM_ADDRESS_1
    assert details["address_line_2"] == CUSTOM_ADDRESS_2
    assert details["address_line_3"] == CUSTOM_ADDRESS_3


def test_system_stamp_is_backend_derived_dated_and_png_is_real() -> None:
    document = _document()
    fixed_issue_time = datetime(2026, 9, 9, 15, 30, tzinfo=timezone.utc)
    stamp = build_system_stamp(document, issued_at=fixed_issue_time)
    assert stamp.brand == "GUARDRISK"
    assert stamp.issuer == "SYSTEM VERIFIED"
    assert stamp.issued_date == "09 SEP 2026"
    assert stamp.reference == "GR-12345678-V4"
    assert len(stamp.content_hash) == 10
    assert stamp.content_hash.isalnum()
    assert USER_CONTROLLED_STAMP not in stamp.reference
    assert not hasattr(stamp, "legal_name")
    assert not hasattr(stamp, "status")

    png = render_system_stamp_png(document)
    assert png.startswith(b"\x89PNG\r\n\x1a\n")
    assert len(png) > 5_000


def test_pdf_matches_reference_and_uses_backend_system_stamp() -> None:
    document = _document()
    expected_stamp = build_system_stamp(document)
    payload = export_pdf(document)
    assert payload.startswith(b"%PDF")
    reader = PdfReader(BytesIO(payload))
    assert len(reader.pages) >= 2
    assert reader.metadata is not None
    assert reader.metadata.get("/Creator") == "gRisk Backend Document Service"
    assert reader.metadata.get("/Producer") == "gRisk Backend Document Service"

    for page in reader.pages:
        text = page.extract_text() or ""
        assert "GUARDRISK" in text
        assert "DOCUMENT STUDIO" not in text
        assert "Code:" not in text
        assert "GRK-" not in text
        assert CUSTOM_TAGLINE.upper() in text
        assert CUSTOM_DATE in text
        assert CUSTOM_RECIPIENT in text
        assert CUSTOM_COMPANY in text
        assert CUSTOM_SUBJECT in text
        assert CUSTOM_ADDRESS_1 in text
        assert CUSTOM_ADDRESS_2 in text
        assert CUSTOM_ADDRESS_3 in text
        assert "+266 2232 2537" in text
        assert "+266 6272 0488" in text
        assert "custom@guardrisk.co.ls" in text
        assert "GUARDRISK HEALTH" in text
        assert "LOW COST MEDICAL AID" in text

    complete_text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "First page body content." in complete_text
    assert "Second page body content." in complete_text
    assert CUSTOM_CLOSING_1 in complete_text
    assert CUSTOM_CLOSING_2 in complete_text
    assert CUSTOM_SIGNATURE in complete_text
    assert CUSTOM_SIGNER in complete_text
    assert CUSTOM_TITLE in complete_text
    assert "INSURANCE BROKERS" not in complete_text
    assert "OFFICIAL" not in complete_text
    assert "SYSTEM VERIFIED" in complete_text
    assert f"DATE  {expected_stamp.issued_date}" in complete_text
    assert "GR-12345678-V4" in complete_text
    assert "VERIFY " in complete_text
    assert USER_CONTROLLED_STAMP not in complete_text


def test_docx_matches_reference_and_embeds_backend_system_stamp_image() -> None:
    payload = export_docx(_document())
    assert payload.startswith(b"PK")
    word = WordDocument(BytesIO(payload))
    assert word.sections

    for section in word.sections:
        header_text = _container_text(section.header)
        footer_text = _container_text(section.footer)
        assert "GUARDRISK" in header_text
        assert "DOCUMENT STUDIO" not in header_text
        assert "Code:" not in header_text
        assert "GRK-" not in header_text
        assert CUSTOM_TAGLINE.upper() in header_text
        assert CUSTOM_DATE in header_text
        assert CUSTOM_RECIPIENT in header_text
        assert CUSTOM_COMPANY in header_text
        assert CUSTOM_SUBJECT in header_text
        assert CUSTOM_ADDRESS_1 in header_text
        assert CUSTOM_ADDRESS_2 in header_text
        assert CUSTOM_ADDRESS_3 in header_text
        assert CUSTOM_ADDRESS_1 in footer_text
        assert "+266 2232 2537" in footer_text
        assert "+266 6272 0488" in footer_text
        assert "custom@guardrisk.co.ls" in footer_text
        assert "GUARDRISK HEALTH" in footer_text
        assert "LOW COST MEDICAL AID" in footer_text

    body_text = _document_body_text(word)
    assert "First page body content." in body_text
    assert "Second page body content." in body_text
    assert CUSTOM_CLOSING_1 in body_text
    assert CUSTOM_CLOSING_2 in body_text
    assert CUSTOM_SIGNATURE in body_text
    assert CUSTOM_SIGNER in body_text
    assert CUSTOM_TITLE in body_text
    assert USER_CONTROLLED_STAMP not in body_text
    assert len(word.inline_shapes) >= 1
    assert word.core_properties.last_modified_by == "gRisk Backend Document Service"
    assert "system stamp" in (word.core_properties.keywords or "")
