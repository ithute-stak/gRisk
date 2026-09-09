import uuid
from io import BytesIO

from docx import Document as WordDocument
from pypdf import PdfReader

from app.models.studio import StudioDocument
from app.services.studio_export import (
    export_docx,
    export_pdf,
    plain_text_from_html,
    sanitize_document_html,
)


def _document() -> StudioDocument:
    html = """
    <h1 style="text-align:center;color:#3D2A1D">Executive heading</h1>
    <p><strong>Important</strong> Guardrisk document text.</p>
    <ul><li>Item one</li><li>Item two</li></ul>
    <table><tbody><tr><th>Cover</th><th>Premium</th></tr><tr><td>Property</td><td>M 1,250.00</td></tr></tbody></table>
    <div data-signature-field="true" data-field-id="recipient-signature"
         data-field-type="signature" data-label="Recipient signature"
         data-assigned-to="Recipient" data-required="true" data-width="100"
         data-height="62" data-placeholder="Sign here"></div>
    <div data-page-break="true"></div>
    <p>Second page content.</p>
    """
    return StudioDocument(
        id=uuid.uuid4(),
        owner_user_id=uuid.uuid4(),
        title="Structured Export Test",
        template_key="formal_letter",
        style_key="guardrisk_orange",
        content_json={"type": "doc", "content": []},
        html_content=html,
        plain_text=plain_text_from_html(html),
        settings={
            "page_size": "a4",
            "orientation": "portrait",
            "margin_top_mm": 20,
            "margin_right_mm": 20,
            "margin_bottom_mm": 20,
            "margin_left_mm": 20,
            "brand_header": True,
            "default_font_family": "Arial",
            "default_font_size_pt": 11,
            "default_line_height_percent": 115,
        },
        version=3,
    )


def test_document_html_sanitizer_preserves_fields_but_removes_active_content() -> None:
    source = (
        '<p onclick="alert(1)" style="color:#F47A20">Safe</p>'
        '<script>alert("bad")</script>'
        '<div data-signature-field="true" data-field-id="sig-1" data-field-type="signature" '
        'data-label="Signature" data-assigned-to="Recipient" data-required="true"></div>'
    )
    clean = sanitize_document_html(source)
    assert "<script" not in clean
    assert "onclick" not in clean
    assert "color:#F47A20" in clean
    assert 'data-signature-field="true"' in clean
    assert "[Signature]" in plain_text_from_html(clean)


def test_document_studio_exports_real_docx_with_tables_and_fields() -> None:
    payload = export_docx(_document())
    assert payload.startswith(b"PK")
    word = WordDocument(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    table_text = "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
    assert "Executive heading" in text
    assert "Second page content" in text
    assert "Cover" in table_text
    assert "Recipient signature" in table_text
    assert len(word.tables) >= 2


def test_document_studio_exports_rich_pdf() -> None:
    payload = export_pdf(_document())
    assert payload.startswith(b"%PDF")
    reader = PdfReader(BytesIO(payload))
    assert len(reader.pages) >= 2
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Executive heading" in text
    assert "Recipient signature" in text
    assert "Second page content" in text
